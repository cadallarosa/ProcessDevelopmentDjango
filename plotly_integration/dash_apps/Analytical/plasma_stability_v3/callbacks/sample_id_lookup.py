"""
Sample ID Lookup Callbacks

Handles the Sample ID Lookup workflow:
- Upload template with Sample IDs
- Look up Result IDs from database
- Flag duplicates
- Validate Sample ID/Result ID pairs
- Export enhanced Excel with validation info
"""

import pandas as pd
import base64
import io
from dash import Input, Output, State, html, dcc
from dash.exceptions import PreventUpdate
import dash_bootstrap_components as dbc
from datetime import datetime

from ..utils.sample_lookup import process_sample_lookup


def register_callbacks(app):
    """Register all Sample ID Lookup callbacks."""

    @app.callback(
        [
            Output('sample-lookup-results-store', 'data'),
            Output('sample-lookup-status', 'children'),
            Output('sample-lookup-download-section', 'style')
        ],
        Input('sample-lookup-upload', 'contents'),
        State('sample-lookup-upload', 'filename')
    )
    def process_sample_lookup_template(contents, filename):
        """
        Process uploaded template and perform Sample ID lookups.
        """
        if not contents:
            raise PreventUpdate

        try:
            # Decode and read Excel file
            content_type, content_string = contents.split(',')
            decoded = base64.b64decode(content_string)

            # Read the Sample Configuration sheet (header is on row 2, index 1)
            df = pd.read_excel(io.BytesIO(decoded), sheet_name='Sample Configuration', header=1)

            # Clean up the dataframe
            df = df.dropna(how='all')

            # Filter out example rows (those with '#' in Result ID)
            df = df[~df['Result ID'].astype(str).str.contains('#', na=False)]

            # Process each row
            results = []
            for idx, row in df.iterrows():
                sample_id = row.get('Sample ID', None)
                result_id = row.get('Result ID', None)

                # Convert to string and handle NaN
                if pd.isna(sample_id):
                    sample_id = None
                else:
                    sample_id = str(sample_id).strip()

                if pd.isna(result_id) or str(result_id).strip() == '':
                    result_id = None
                else:
                    try:
                        result_id = int(float(result_id))
                    except (ValueError, TypeError):
                        result_id = str(result_id).strip()

                # Perform lookup
                lookup_result = process_sample_lookup(sample_id, result_id)

                # Combine with original row data
                result_row = {
                    'Instrument': row.get('Instrument', ''),
                    'Result ID': result_id if result_id else '',
                    'Sample ID': sample_id if sample_id else '',
                    'Molecule ID': row.get('Molecule ID', ''),
                    'Matrix': row.get('Matrix', ''),
                    'Day': row.get('Day', ''),
                    'Peak RT Type': row.get('Peak RT Type', 0),
                    'Peak RT': row.get('Peak RT', ''),
                    # Lookup results
                    'Found Result IDs': '|'.join(lookup_result['found_result_ids']),
                    'Lookup Status': lookup_result['status'],
                    'Duplicate Flag': lookup_result['duplicate_flag'],
                    'Duplicate Count': lookup_result['duplicate_count'],
                    'Validation Notes': lookup_result['validation_notes'],
                    'Date Acquired': lookup_result['date_acquired'],
                    'Acquired By': lookup_result['acquired_by']
                }
                results.append(result_row)

            # Create results DataFrame
            results_df = pd.DataFrame(results)

            # Calculate statistics
            total_samples = len(results_df)
            found_count = len(results_df[results_df['Lookup Status'].isin(['Found', 'Match Confirmed'])])
            duplicate_count = len(results_df[results_df['Duplicate Flag'] == 'Yes'])
            not_found_count = len(results_df[results_df['Lookup Status'] == 'Not Found'])
            error_count = len(results_df[results_df['Lookup Status'] == 'Validation Error'])

            # Create status display
            status_display = dbc.Card([
                dbc.CardBody([
                    html.H5("Processing Summary", className="mb-3"),
                    dbc.Row([
                        dbc.Col([
                            html.Div([
                                html.I(className="fas fa-check-circle text-success me-2"),
                                html.Span(f"{total_samples} samples processed", className="fw-bold")
                            ], className="mb-2")
                        ]),
                    ]),
                    dbc.Row([
                        dbc.Col([
                            html.Div([
                                html.I(className="fas fa-check text-success me-2"),
                                html.Span(f"{found_count} found successfully")
                            ], className="mb-1")
                        ], width=6),
                        dbc.Col([
                            html.Div([
                                html.I(className="fas fa-exclamation-triangle text-warning me-2"),
                                html.Span(f"{duplicate_count} with duplicates")
                            ], className="mb-1")
                        ], width=6),
                    ]),
                    dbc.Row([
                        dbc.Col([
                            html.Div([
                                html.I(className="fas fa-times text-danger me-2"),
                                html.Span(f"{not_found_count} not found")
                            ], className="mb-1")
                        ], width=6),
                        dbc.Col([
                            html.Div([
                                html.I(className="fas fa-exclamation-circle text-danger me-2"),
                                html.Span(f"{error_count} validation errors")
                            ], className="mb-1")
                        ], width=6),
                    ]),
                ])
            ], className="mt-3 mb-3")

            # Store results as JSON
            results_json = results_df.to_dict('records')

            # Show download section
            download_style = {'display': 'block'}

            return results_json, status_display, download_style

        except Exception as e:
            error_display = dbc.Alert([
                html.H5("Error Processing Template", className="alert-heading"),
                html.P(f"An error occurred while processing the template: {str(e)}"),
                html.Hr(),
                html.P("Please ensure you're uploading the correct template format.", className="mb-0")
            ], color="danger")

            return None, error_display, {'display': 'none'}

    @app.callback(
        Output('sample-lookup-download', 'data'),
        Input('sample-lookup-download-btn', 'n_clicks'),
        State('sample-lookup-results-store', 'data'),
        State('sample-lookup-upload', 'filename'),
        prevent_initial_call=True
    )
    def download_enhanced_template(n_clicks, results_data, original_filename):
        """
        Generate and download enhanced Excel template with validation info.
        """
        if not n_clicks or not results_data:
            raise PreventUpdate

        try:
            # Convert results to DataFrame
            df = pd.DataFrame(results_data)

            # Create Excel writer with multiple sheets
            output = io.BytesIO()
            with pd.ExcelWriter(output, engine='openpyxl') as writer:
                # Write main results
                df.to_excel(writer, sheet_name='Sample Lookup Results', index=False)

                # Get workbook and worksheet for formatting
                workbook = writer.book
                worksheet = writer.sheets['Sample Lookup Results']

                # Apply conditional formatting
                from openpyxl.styles import PatternFill, Font

                # Define styles
                success_fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
                warning_fill = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')
                error_fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')

                # Get column indices
                status_col = None
                duplicate_col = None
                for idx, col in enumerate(df.columns, start=1):
                    if col == 'Lookup Status':
                        status_col = idx
                    elif col == 'Duplicate Flag':
                        duplicate_col = idx

                # Apply formatting to each row
                for row_idx, row in enumerate(df.itertuples(), start=2):  # Start at 2 (skip header)
                    status = row._asdict().get('Lookup Status', '')
                    duplicate_flag = row._asdict().get('Duplicate Flag', '')

                    # Color entire row based on status
                    fill = None
                    if status in ['Match Confirmed', 'Found']:
                        fill = success_fill
                    elif status == 'Multiple Matches' or duplicate_flag == 'Yes':
                        fill = warning_fill
                    elif status in ['Not Found', 'Validation Error']:
                        fill = error_fill

                    if fill:
                        for col_idx in range(1, len(df.columns) + 1):
                            worksheet.cell(row=row_idx, column=col_idx).fill = fill

                # Adjust column widths
                for col_idx, col in enumerate(df.columns, start=1):
                    max_length = max(
                        df[col].astype(str).apply(len).max(),
                        len(str(col))
                    )
                    adjusted_width = min(max_length + 2, 50)
                    worksheet.column_dimensions[worksheet.cell(row=1, column=col_idx).column_letter].width = adjusted_width

                # Create summary sheet
                summary_data = {
                    'Metric': [
                        'Total Samples',
                        'Successfully Matched',
                        'Duplicates Found',
                        'Not Found',
                        'Validation Errors'
                    ],
                    'Count': [
                        len(df),
                        len(df[df['Lookup Status'].isin(['Found', 'Match Confirmed'])]),
                        len(df[df['Duplicate Flag'] == 'Yes']),
                        len(df[df['Lookup Status'] == 'Not Found']),
                        len(df[df['Lookup Status'] == 'Validation Error'])
                    ]
                }
                summary_df = pd.DataFrame(summary_data)
                summary_df.to_excel(writer, sheet_name='Summary', index=False)

            # Prepare download
            output.seek(0)

            # Generate filename
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            base_name = original_filename.replace('.xlsx', '') if original_filename else 'template'
            download_filename = f"{base_name}_enhanced_{timestamp}.xlsx"

            return dcc.send_bytes(output.getvalue(), download_filename)

        except Exception as e:
            print(f"Error generating Excel download: {e}")
            raise PreventUpdate

    @app.callback(
        Output('sample-lookup-preview-table', 'children'),
        Input('sample-lookup-results-store', 'data')
    )
    def update_preview_table(results_data):
        """
        Display preview of first 10 results.
        """
        if not results_data:
            return html.Div("No data to display", className="text-muted text-center p-3")

        try:
            df = pd.DataFrame(results_data)

            # Show first 10 rows
            preview_df = df.head(10)

            # Select key columns for preview
            preview_cols = [
                'Sample ID', 'Result ID', 'Found Result IDs',
                'Lookup Status', 'Duplicate Flag', 'Validation Notes'
            ]
            preview_df = preview_df[preview_cols]

            # Create table
            table = dbc.Table.from_dataframe(
                preview_df,
                striped=True,
                bordered=True,
                hover=True,
                size='sm',
                className='mt-2'
            )

            return html.Div([
                html.H6(f"Preview (showing first 10 of {len(df)} rows)"),
                table
            ])

        except Exception as e:
            return html.Div(f"Error displaying preview: {str(e)}", className="text-danger")
